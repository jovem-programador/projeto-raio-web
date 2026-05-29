"""
Developer: Anderson Marley
GitHub: jovem-programador
Date: 02/03/2026
"""

import os
import re
import shutil
import subprocess
import unicodedata
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
import ezdxf
import pandas as pd

# ================= CONFIGURAÇÕES DE CAMINHO =================
PATH_ODA = r"C:\Program Files\ODA\ODAFileConverter 26.12.0\ODAFileConverter.exe"
PASTA_DWG_ORIGEM = r"C:\Users\anderson.marley\Documents\Projeta\5 - Projeto 'O RAIO'\Projetos_DWG"
PASTA_DXF_TEMP = r"C:\Users\anderson.marley\Documents\Projeta\5 - Projeto 'O RAIO'\DXF_Temporario"
ARQUIVO_EXCEL_SAIDA = r"C:\Users\anderson.marley\Documents\Projeta\5 - Projeto 'O RAIO'\Excel Raio\Extração_Carimbos_Vale.xlsx"

DXF_VERSION = "ACAD2018"
CONVERTER_DWG_PARA_DXF = True

CAMPOS_ORDEM = [
    "CLASSIFICAÇÃO",
    "PROJETO",
    "N_PROJ_SE",
    "N_CONTRATO",
    "FASE_DO_PROJETO",
    "ÁREA_E/OU_SUBÁREA",
    "TÍTULO_DO_DESENHO",
    "SUBTÍTULO_1_DO_DESENHO",
    "SUBTÍTULO_2_DO_DESENHO",
    "ESCALA",
    "NÚMERO_DA_CONTRATADA",
    "NUMERO_VALE",
    "1a_REV",
    "EMIS_1",
    "DESCRIÇÃO_DA_REVISÃO_1",
    "2a_REV",
    "EMIS_2",
    "DESCRIÇÃO_DA_REVISÃO_2",
]

ALIASES = {
    "CLASSIFICAÇÃO": ["CLASSIFICAÇÃO", "CLASSIFICACAO"],
    "PROJETO": ["PROJETO", "PROJ.", "PROJ:"],

    "N_PROJ_SE": [
        "Nº DO PROJETO OU SE", "N° DO PROJETO OU SE", "NR DO PROJETO OU SE", "N DO PROJETO OU SE",
        "Nº DO PROJETO", "N° DO PROJETO", "NR DO PROJETO", "N DO PROJETO",
        "N_PROJ_SE",
    ],

    "N_CONTRATO": [
    "Nº DO CONTRATO", "N° DO CONTRATO", "NR DO CONTRATO", "N DO CONTRATO",
    "Nº CONTRATO", "N° CONTRATO", "NR CONTRATO", "N CONTRATO",
    "CONTRATO", "CONTRATO Nº", "CONTRATO N°", "CONTRATO NR",
    "N_CONTRATO",

    "SE", "Nº SE", "N° SE", "NR SE", "N SE",
    ],

    "FASE_DO_PROJETO": [
        "FASE DO PROJETO", "FASE_DO_PROJETO", "FASE", "FASE.", "FASE:",
    ],

    "ÁREA_E/OU_SUBÁREA": [
        "ÁREA E/OU SUBÁREA", "AREA E/OU SUBAREA",
        "ÁREA", "AREA", "SUBÁREA", "SUBAREA",
    ],

    "TÍTULO_DO_DESENHO": [
    "TÍTULO DO DESENHO", "TITULO DO DESENHO",
    "TÍTULO", "TITULO",
    "TIT.", "TIT:", "TIT", "TÍT.", "TÍT:",
    ],

    "SUBTÍTULO_1_DO_DESENHO": [
        "SUBTÍTULO 1 DO DESENHO", "SUBTITULO 1 DO DESENHO", "SUBTÍTULO 1", "SUBTITULO 1",
        "SUB1", "SUB1.", "SUB1:", "SUBT1", "SUBT1.", "SUBT1:", "SUBTITULO1", "SUBTÍTULO1",
        "SUBTIT.1", "SUBTÍT.1", "SUBTIT. 1", "SUBTÍT. 1",
    ],

    "SUBTÍTULO_2_DO_DESENHO": [
        "SUBTÍTULO 2 DO DESENHO", "SUBTITULO 2 DO DESENHO", "SUBTÍTULO 2", "SUBTITULO 2",
        "SUB2", "SUB2.", "SUB2:", "SUBT2", "SUBT2.", "SUBT2:", "SUBTITULO2", "SUBTÍTULO2",
        "SUBTIT.2", "SUBTÍT.2", "SUBTIT. 2", "SUBTÍT. 2",
    ],

    "ESCALA": ["ESCALA", "ESC.", "ESC:"],

    "NÚMERO_DA_CONTRATADA": [
        "Nº CONTRATADA", "N° CONTRATADA", "N CONTRATADA", "Nº DA CONTRATADA", "N° DA CONTRATADA",
        "DG-PACG", "NÚMERO_DA_CONTRATADA",
    ],

    "NUMERO_VALE": [
        "Nº VALE", "N° VALE", "N VALE",
        "Nº DO DESENHO VALE", "N° DO DESENHO VALE",
        "VALE", "NUMERO_VALE",
    ],

}

# OBS: estes aliases são usados apenas no modo "por texto".
# Para ATTRIB, a lógica principal é contagem por ocorrência.
ALIAS_EMIS = ["EMIS.", "EMIS", "EMISSÃO", "EMISSAO"]
ALIAS_DESC = ["DESCRIÇÃO", "DESCRICAO", "DESCRIÇÃO DA REVISÃO", "DESCRICAO DA REVISAO", "DESC"]


# ================= UTIL =================
def _near_value_for_key(items: list[dict], key_item: dict, y_tol: float = 2.0) -> str:
    """
    Procura o valor mais provável para uma chave que aparece sozinha.
    Estratégia:
      1) texto mais próximo à DIREITA na mesma linha (|dy| <= y_tol)
      2) se não achar, texto mais próximo ABAIXO (mesmo x aproximado)
    """
    x0, y0 = key_item.get("x"), key_item.get("y")
    if x0 is None or y0 is None:
        return ""

    candidates = []
    for it in items:
        txt = (it.get("text") or "").strip()
        if not txt:
            continue
        x, y = it.get("x"), it.get("y")
        if x is None or y is None:
            continue

        # 1) direita mesma linha
        if x > x0 and abs(y - y0) <= y_tol:
            candidates.append(("right", x - x0, abs(y - y0), txt))

    if candidates:
        candidates.sort(key=lambda t: (t[1], t[2]))  # menor dx, depois menor dy
        return candidates[0][3].strip()

    # 2) abaixo (fallback)
    below = []
    for it in items:
        txt = (it.get("text") or "").strip()
        if not txt:
            continue
        x, y = it.get("x"), it.get("y")
        if x is None or y is None:
            continue
        if y < y0 and abs(x - x0) <= 5.0:  # tolerância em X
            below.append((y0 - y, abs(x - x0), txt))
    if below:
        below.sort(key=lambda t: (t[0], t[1]))  # menor dy para baixo, depois menor dx
        return below[0][2].strip()

    return ""

def normalize(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip().upper()


def normalize_key(s: str) -> str:
    """Normalização tolerante para TAGs/aliases (remove acentos, trata '_' como espaço)."""
    s = (s or "").strip().upper().replace("_", " ")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = re.sub(r"\s+", " ", s).strip()
    return s

def clean_mtext(s: str) -> str:
    if not s:
        return ""
    s = s.replace(r"\P", "\n")
    s = re.sub(r"{\\.*?;|}", "", s)            # remove grupos {...}
    s = re.sub(r"\\[A-Za-z0-9]+", "", s)       # remove comandos \H \W etc (básico)
    s = s.replace("\\~", " ")
    return s.strip()

def pick_value_after_key(full_text: str, key: str) -> str:
    """
    Extrai valor após a chave, com separadores comuns.
    Ex: 'Nº DO CONTRATO 5500081543' -> '5500081543'
    """
    up = normalize(full_text)
    ku = normalize(key)
    m = re.search(re.escape(ku) + r"[\s:=-]*([A-Z0-9\.\-/]+.*)$", up)
    return m.group(1).strip() if m else ""


def pick_nearest_in_window(
    items: list[dict],
    x_ref: float,
    y_ref: float,
    x_tol: float = 60.0,
    y_tol: float = 2.0,
    kind_prefer: tuple[str, ...] = ("ATTRIB", "TEXT", "MTEXT"),
) -> str:
    """
    Extrai o texto mais próximo do ponto (x_ref, y_ref) dentro de uma janela (x_tol/y_tol).

    Útil como fallback quando:
      - a TAG/alias do atributo varia entre carimbos
      - o campo vem como TEXT/MTEXT sem TAG
      - queremos buscar por coordenada (carimbo padronizado)

    Observação:
      - use y_tol baixo quando existirem campos muito próximos na vertical
        (ex.: N_PROJ_SE e SE no carimbo VALE).
    """
    x_min, x_max = x_ref - x_tol, x_ref + x_tol
    y_min, y_max = y_ref - y_tol, y_ref + y_tol

    candidates: list[tuple[float, str, str]] = []
    for it in items:
        x = it.get("x")
        y = it.get("y")
        txt = (it.get("text") or "").strip()
        if x is None or y is None or not txt:
            continue

        if x_min <= x <= x_max and y_min <= y <= y_max:
            dist = ((x - x_ref) ** 2 + (y - y_ref) ** 2) ** 0.5
            candidates.append((dist, (it.get("kind") or ""), txt))

    if not candidates:
        return ""

    # Preferência por origem do texto (ATTRIB > TEXT > MTEXT)
    for k in kind_prefer:
        same_kind = [c for c in candidates if c[1] == k]
        if same_kind:
            same_kind.sort(key=lambda t: t[0])
            return same_kind[0][2].strip()

    # Fallback: mais próximo geral
    candidates.sort(key=lambda t: t[0])
    return candidates[0][2].strip()


def find_first_by_tag(items: list[dict], tag_name: str) -> dict | None:
    """Retorna o primeiro item ATTRIB/TEXT que tenha TAG exatamente (após normalização) e VAL não vazio."""
    target = normalize_key(tag_name)
    for it in items:
        if normalize_key(it.get("tag") or "") == target and (it.get("text") or "").strip():
            return it
    return None

def pick_nearest_relative(
    items: list[dict],
    x_ref_base: float,
    y_ref_base: float,
    x_anchor_base: float,
    y_anchor_base: float,
    x_anchor_real: float,
    y_anchor_real: float,
    x_tol: float = 120.0,
    y_tol: float = 4.0,
) -> str:
    """Aplica deslocamento (dx,dy) a partir da âncora e busca na janela."""
    dx = x_anchor_real - x_anchor_base
    dy = y_anchor_real - y_anchor_base
    return pick_nearest_in_window(
        items,
        x_ref=x_ref_base + dx,
        y_ref=y_ref_base + dy,
        x_tol=x_tol,
        y_tol=y_tol,
    )





# ================= EXTRAÇÃO (ATTRIB + TEXT + MTEXT, TODOS OS LAYOUTS) =================
def extract_all_texts(dxf_path: str) -> list[dict]:
    try:
        doc = ezdxf.readfile(dxf_path)
    except Exception as e:
        print(f"Erro ao ler DXF {dxf_path}: {e}")
        return []

    out = []
    for layout in doc.layouts:
        layout_name = layout.name

        for entity in layout:
            dxft = entity.dxftype()

            if dxft == "INSERT":
                for attr in getattr(entity, "attribs", []):
                    tag = (attr.dxf.tag or "").strip().upper()
                    text = (attr.dxf.text or "").strip()
                    ins = getattr(attr.dxf, "insert", None)
                    x, y = (float(ins.x), float(ins.y)) if ins else (None, None)
                    out.append({"layout": layout_name, "tag": tag, "text": text, "x": x, "y": y, "kind": "ATTRIB"})

            elif dxft == "TEXT":
                text = (entity.dxf.text or "").strip()
                ins = getattr(entity.dxf, "insert", None)
                x, y = (float(ins.x), float(ins.y)) if ins else (None, None)
                out.append({"layout": layout_name, "tag": "TEXT_ENTITY", "text": text, "x": x, "y": y, "kind": "TEXT"})

            elif dxft == "MTEXT":
                text = clean_mtext(getattr(entity, "text", "") or "")
                text = re.sub(
                    r"\\P|\\f[^;]*;|\\H[^;]*;|\\S[^;]*;|\\Q[^;]*;|\\W[^;]*;|\\A[^;]*;|\\L|\\l|\\O|\\o|{|}|\\~",
                    " ",
                    text,
                ).strip()
                ins = getattr(entity.dxf, "insert", None)
                x, y = (float(ins.x), float(ins.y)) if ins else (None, None)
                out.append({"layout": layout_name, "tag": "TEXT_ENTITY", "text": text, "x": x, "y": y, "kind": "MTEXT"})

    return out

# ================= PREENCHIMENTO (TAG + TEXTO “CHAVE VALOR”) =================
def fill_by_alias_tags(items: list[dict], carimbo: dict):
    """
    Preenche por TAG (quando o carimbo vem em ATTRIB e a tag é igual ao alias).
    Ex: tag == 'PROJETO' -> valor direto.
    """
    for item in items:
        tag = normalize_key(item.get("tag") or "")
        val = (item.get("text") or "").strip()
        if not val:
            continue

        for campo, aliases in ALIASES.items():
            for a in aliases:
                if tag == normalize_key(a):
                    if not carimbo.get(campo):
                        carimbo[campo] = val

def fill_by_key_value_text(items: list[dict], carimbo: dict):
    """
    Preenche por texto “chave + valor” (bom para TEXT/MTEXT).
    Ex: 'Nº DO CONTRATO: 5500...' -> extrai 5500...
    """
    texts = [(it.get("text") or "").strip() for it in items if (it.get("text") or "").strip()]
    for campo in [
        "CLASSIFICAÇÃO",
        "PROJETO",
        "N_PROJ_SE",
        "N_CONTRATO",
        "FASE_DO_PROJETO",
        "ÁREA_E/OU_SUBÁREA",
        "TÍTULO_DO_DESENHO",
        "SUBTÍTULO_1_DO_DESENHO",
        "SUBTÍTULO_2_DO_DESENHO",
        "ESCALA",
        "NÚMERO_DA_CONTRATADA",
        "NUMERO_VALE",
    ]:
        if carimbo.get(campo):
            continue
        for k in sorted(ALIASES.get(campo, []), key=lambda x: len(normalize(x)), reverse=True):
            ku = normalize(k)
            for t in texts:
                if normalize(t).startswith(ku):
                    v = pick_value_after_key(t, k)
                    if v:
                        carimbo[campo] = v
                        break
            if carimbo.get(campo):
                break


# ================= ÁREA/SUBÁREA (CORREÇÃO) =================
def processar_area_subarea_por_tag(items: list[dict], carimbo: dict):
    """
    Corrige casos em que o carimbo tem ÁREA e SUBÁREA em ATTRIBs separados.
    Monta o campo ÁREA_E/OU_SUBÁREA no formato: 'ÁREA: X | SUBÁREA: Y'
    """
    area = ""
    sub = ""

    for item in items:
        tag = (item.get("tag") or "").strip().upper()
        val = (item.get("text") or "").strip()
        if not val:
            continue

        # pega separado (muito comum em carimbo VALE)
        if tag in ("ÁREA", "AREA", "ÁREA:", "AREA:"):
            area = val
        elif tag in ("SUBÁREA", "SUBAREA", "SUBÁREA:", "SUBAREA:"):
            sub = val

    # Se encontrou algo separado, compõe (e sobrescreve o genérico)
    if area or sub:
        if area and sub:
            carimbo["ÁREA_E/OU_SUBÁREA"] = f"ÁREA: {area} | SUBÁREA: {sub}"
        elif area:
            carimbo["ÁREA_E/OU_SUBÁREA"] = f"ÁREA: {area}"
        else:
            carimbo["ÁREA_E/OU_SUBÁREA"] = f"SUBÁREA: {sub}"


# ================= REVISÕES (CORRIGIDO: SEQUÊNCIA REAL DO CARIMBO) =================
def processar_revisoes_dinamicas_por_tag(items: list[dict], carimbo: dict):
    """
    Este é o ponto que corrige seu caso do print:
      1A_REV / EMIS. / DESCRIÇÃO_DA_REVISÃO / ... / 2A_REV / EMIS. / DESCRIÇÃO_DA_REVISÃO / ...

    Como EMIS. e DESCRIÇÃO_DA_REVISÃO não possuem índice na TAG,
    a única forma robusta é contar ocorrências (1ª = _1, 2ª = _2).
    """
    contador_emis = 0
    contador_desc = 0

    for item in items:
        tag = (item.get("tag") or "").strip().upper()
        val = (item.get("text") or "").strip()
        if not val:
            continue

        # 1) 1A_REV / 2A_REV / 3A_REV...
        m_rev = re.search(r"(\d+)\s*[Aª]?\s*_?\s*REV", tag)
        if m_rev:
            key = f"{m_rev.group(1)}a_REV"
            if key in carimbo and not carimbo[key]:
                carimbo[key] = val
            continue

        # 2) EMIS / EMIS. / EMISSÃO...
        if tag.startswith("EMIS"):
            contador_emis += 1
            key = f"EMIS_{contador_emis}"
            if key in carimbo and not carimbo[key]:
                carimbo[key] = val
            continue

        # 3) DESCRIÇÃO_DA_REVISÃO / DESCRI...
        if "DESCRI" in tag:
            contador_desc += 1
            key = f"DESCRIÇÃO_DA_REVISÃO_{contador_desc}"
            if key in carimbo and not carimbo[key]:
                carimbo[key] = val
            continue

        # 4) DATA — (se quiser salvar depois, dá pra criar DATA_1/DATA_2)
        # if tag == "DATA": ...


def processar_revisoes_por_sequencia_texto(items: list[dict], carimbo: dict):
    """
    Complemento: caso a revisão venha em TEXT/MTEXT (chave + valor),
    tenta pegar:
      - "1a REV: A" / "2a REV: 0"
      - "EMIS.: 23/08/25" (ou similar) -> EMIS_1/EMIS_2 por ocorrência
      - "DESCRIÇÃO: ..." -> DESCRIÇÃO_DA_REVISÃO_1/2 por ocorrência
    """
    texts = [(it.get("text") or "").strip() for it in items if (it.get("text") or "").strip()]

    # REV por texto
    if not carimbo.get("1a_REV"):
        for t in texts:
            m = re.search(r"(1\s*[ªAa]?\s*REV)[\s:=-]*(.+)$", normalize(t))
            if m:
                carimbo["1a_REV"] = m.group(2).strip()
                break

    if not carimbo.get("2a_REV"):
        for t in texts:
            m = re.search(r"(2\s*[ªAa]?\s*REV)[\s:=-]*(.+)$", normalize(t))
            if m:
                carimbo["2a_REV"] = m.group(2).strip()
                break

    # EMIS e DESCRIÇÃO por ocorrência
    emis_vals = []
    desc_vals = []

    for t in texts:
        up = normalize(t)

        if any(normalize(a) in up for a in ALIAS_EMIS):
            v = ""
            for a in ALIAS_EMIS:
                if normalize(a) in up:
                    v = pick_value_after_key(t, a)
                    if v:
                        break
            if v:
                emis_vals.append(v)

        if any(normalize(a) in up for a in ALIAS_DESC):
            v = ""
            for a in ALIAS_DESC:
                if normalize(a) in up:
                    v = pick_value_after_key(t, a)
                    if v:
                        break
            if v:
                desc_vals.append(v)

    if emis_vals and not carimbo.get("EMIS_1"):
        carimbo["EMIS_1"] = emis_vals[0]
    if len(emis_vals) > 1 and not carimbo.get("EMIS_2"):
        carimbo["EMIS_2"] = emis_vals[1]

    if desc_vals and not carimbo.get("DESCRIÇÃO_DA_REVISÃO_1"):
        carimbo["DESCRIÇÃO_DA_REVISÃO_1"] = desc_vals[0]
    if len(desc_vals) > 1 and not carimbo.get("DESCRIÇÃO_DA_REVISÃO_2"):
        carimbo["DESCRIÇÃO_DA_REVISÃO_2"] = desc_vals[1]


def _is_placeholder(v: str) -> bool:
    v = normalize(v)
    if not v:
        return True
    # placeholders típicos de template
    if re.fullmatch(r"[XN]+", v):  # XXXX / NNNNN
        return True
    if "TITULO DO DESENHO" in v or "SUBTITULO" in v:
        return True
    if "AREA E/OU SUBAREA" in v or "ÁREA E/OU SUBÁREA" in v:
        return True
    if v in {"FASE DO PROJETO", "PROJETO"}:
        return True
    return False


def _score_carimbo(carimbo: dict) -> int:
    # conta quantos campos relevantes têm valor "real"
    campos_relevantes = [
        "CLASSIFICAÇÃO",
        "PROJETO",
        "N_PROJ_SE",
        "N_CONTRATO",
        "FASE_DO_PROJETO",
        "ÁREA_E/OU_SUBÁREA",
        "TÍTULO_DO_DESENHO",
        "SUBTÍTULO_1_DO_DESENHO",
        "SUBTÍTULO_2_DO_DESENHO",
        "NUMERO_VALE",
        "1a_REV",
    ]
    score = 0
    for c in campos_relevantes:
        if not _is_placeholder(carimbo.get(c, "")):
            score += 1
    return score


def _extrair_carimbo_de_items(items: list[dict], dxf_path: str, x_tol: float, y_tol: float) -> dict | None:
    if not items:
        return None

    carimbo = {c: "" for c in CAMPOS_ORDEM}
    carimbo["Nome_Arquivo"] = os.path.basename(dxf_path).replace(".dxf", ".dwg")

    # 1) TAG
    fill_by_alias_tags(items, carimbo)
    # 2) TEXTO
    fill_by_key_value_text(items, carimbo)
    # 3) composições
    processar_area_subarea_por_tag(items, carimbo)
    # 4) revisões
    processar_revisoes_dinamicas_por_tag(items, carimbo)
    processar_revisoes_por_sequencia_texto(items, carimbo)

    # ===== FALLBACK POR COORDENADAS =====
    ANCHOR_N_PROJ_SE_BASE = (3763.91, 1303.01)
    ANCHOR_SE_BASE = (3763.91, 1313.05)

    if not carimbo.get("N_PROJ_SE"):
        carimbo["N_PROJ_SE"] = pick_nearest_in_window(
            items,
            x_ref=ANCHOR_N_PROJ_SE_BASE[0], y_ref=ANCHOR_N_PROJ_SE_BASE[1],
            x_tol=120.0, y_tol=2.0,
        )

    if not carimbo.get("N_CONTRATO"):
        carimbo["N_CONTRATO"] = pick_nearest_in_window(
            items,
            x_ref=ANCHOR_SE_BASE[0], y_ref=ANCHOR_SE_BASE[1],
            x_tol=120.0, y_tol=2.0,
        )

    anchor_item = find_first_by_tag(items, "N_PROJ_SE")
    anchor_base = ANCHOR_N_PROJ_SE_BASE

    if anchor_item is None:
        anchor_item = find_first_by_tag(items, "SE")
        anchor_base = ANCHOR_SE_BASE

    if anchor_item is not None:
        ax = float(anchor_item["x"])
        ay = float(anchor_item["y"])

        if not carimbo.get("ÁREA_E/OU_SUBÁREA"):
            carimbo["ÁREA_E/OU_SUBÁREA"] = pick_nearest_relative(
                items,
                x_ref_base=3581.14, y_ref_base=1295.59,
                x_anchor_base=anchor_base[0], y_anchor_base=anchor_base[1],
                x_anchor_real=ax, y_anchor_real=ay,
                x_tol=min(220.0, x_tol), y_tol=y_tol,
            )

        if not carimbo.get("TÍTULO_DO_DESENHO"):
            carimbo["TÍTULO_DO_DESENHO"] = pick_nearest_relative(
                items,
                x_ref_base=3581.14, y_ref_base=1287.34,
                x_anchor_base=anchor_base[0], y_anchor_base=anchor_base[1],
                x_anchor_real=ax, y_anchor_real=ay,
                x_tol=x_tol, y_tol=y_tol,
            )

        if not carimbo.get("SUBTÍTULO_1_DO_DESENHO"):
            carimbo["SUBTÍTULO_1_DO_DESENHO"] = pick_nearest_relative(
                items,
                x_ref_base=3581.14, y_ref_base=1279.09,
                x_anchor_base=anchor_base[0], y_anchor_base=anchor_base[1],
                x_anchor_real=ax, y_anchor_real=ay,
                x_tol=x_tol, y_tol=y_tol,
            )

        if not carimbo.get("SUBTÍTULO_2_DO_DESENHO"):
            carimbo["SUBTÍTULO_2_DO_DESENHO"] = pick_nearest_relative(
                items,
                x_ref_base=3581.14, y_ref_base=1270.84,
                x_anchor_base=anchor_base[0], y_anchor_base=anchor_base[1],
                x_anchor_real=ax, y_anchor_real=ay,
                x_tol=x_tol, y_tol=y_tol,
            )

    # validação mínima: se tem ao menos algo útil, devolve
    if carimbo.get("PROJETO") or carimbo.get("NUMERO_VALE") or carimbo.get("1a_REV") or carimbo.get("N_PROJ_SE"):
        return carimbo
    return None


def extrair_carimbo_de_um_dxf(dxf_path: str, x_tol: float = 420.0, y_tol: float = 6.0) -> dict | None:
    items = extract_all_texts(dxf_path)
    if not items:
        return None

    # separa por layout (precisa do campo "layout" que você adicionou no extract_all_texts)
    layouts = sorted({it.get("layout", "") for it in items if it.get("layout")})
    if not layouts:
        # fallback: comportamento antigo
        return _extrair_carimbo_de_items(items, dxf_path, x_tol, y_tol)

    melhor = None
    melhor_score = -1
    melhor_layout = None

    for lname in layouts:
        items_l = [it for it in items if it.get("layout") == lname]
        car = _extrair_carimbo_de_items(items_l, dxf_path, x_tol, y_tol)
        if not car:
            continue
        sc = _score_carimbo(car)

        # desempate: preferir Model quando score igual
        if sc > melhor_score or (sc == melhor_score and lname.upper() == "MODEL"):
            melhor = car
            melhor_score = sc
            melhor_layout = lname

    # opcional: registrar layout escolhido (ajuda debug)
    if melhor is not None:
        melhor["_LAYOUT_ESCOLHIDO"] = melhor_layout or ""

    return melhor


# ================= FUNÇÕES DE AMBIENTE =================

def preparar_pastas():
    if os.path.exists(PASTA_DXF_TEMP):
        shutil.rmtree(PASTA_DXF_TEMP)
    os.makedirs(PASTA_DXF_TEMP, exist_ok=True)

def converter_arquivos():
    if not os.path.exists(PATH_ODA):
        raise FileNotFoundError(f"[ERRO] ODA Converter não localizado em: {PATH_ODA}")

    print("---> Convertendo DWG para DXF via ODA...")
    comando = [PATH_ODA, PASTA_DWG_ORIGEM, PASTA_DXF_TEMP, DXF_VERSION, "DXF", "0", "1"]
    subprocess.run(comando, check=True, shell=True)

def extrair_dados_completos():
    resultados = []
    if not os.path.exists(PASTA_DXF_TEMP):
        return resultados

    arquivos = [f for f in os.listdir(PASTA_DXF_TEMP) if f.lower().endswith(".dxf")]
    for arq in arquivos:
        caminho = os.path.join(PASTA_DXF_TEMP, arq)
        carimbo = extrair_carimbo_de_um_dxf(caminho)
        if carimbo:
            resultados.append(carimbo)

    return resultados

def salvar_excel(dados: list[dict]):
    if not dados:
        print("\n[AVISO] Nenhum dado extraído dos desenhos.")
        return

    df = pd.DataFrame(dados)

    cols = ["Nome_Arquivo"] + [c for c in CAMPOS_ORDEM if c in df.columns]
    df = df[cols].drop_duplicates(subset=["Nome_Arquivo"])

    try:
        Path(os.path.dirname(ARQUIVO_EXCEL_SAIDA)).mkdir(parents=True, exist_ok=True)

        # 1️⃣ Salva primeiro o Excel normal
        df.to_excel(ARQUIVO_EXCEL_SAIDA, index=False)

        # 2️⃣ Abre com openpyxl para transformar em TABELA
        wb = load_workbook(ARQUIVO_EXCEL_SAIDA)
        ws = wb.active

        max_row = ws.max_row
        max_col = ws.max_column

        # Define intervalo da tabela
        from openpyxl.utils import get_column_letter
        last_col_letter = get_column_letter(max_col)
        table_range = f"A1:{last_col_letter}{max_row}"

        tab = Table(displayName="Tabela_Raio", ref=table_range)

        style = TableStyleInfo(
            name="TableStyleMedium2",  # você pode trocar estilo aqui
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        tab.tableStyleInfo = style
        ws.add_table(tab)

        wb.save(ARQUIVO_EXCEL_SAIDA)

        print(f"\n---> SUCESSO! Tabela estruturada gerada: {ARQUIVO_EXCEL_SAIDA}")

    except PermissionError:
        print("\n[ERRO] O arquivo Excel está aberto! Feche-o e execute novamente.")

if __name__ == "__main__":
    preparar_pastas()
    if CONVERTER_DWG_PARA_DXF:
        converter_arquivos()
    dados = extrair_dados_completos()
    salvar_excel(dados)
